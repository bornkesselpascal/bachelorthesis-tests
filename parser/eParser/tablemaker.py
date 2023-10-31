import csv
import math
import os
from multiprocessing import Process
import pandas as pd
import openpyxl
from data_format import format_query
from constants import concurrent_execution, generate_excel


def write_test_table(test_data: list, campaign_name: str, campaign_folder: str, processes: list=None) -> None:
    '''
    Creates a table with all test scenarios and their results. Analyzes the data and adds remarks
    if necessary. Checks the NIC and UDP statistics for losses and add hints. Saves the table as
    a CSV file and as an Excel file (optional).

            Parameters:
                    test_data (list): List of parsed test scenarios
                    campaign_folder (str): Path to the folder where the table should be saved
                    processes (list): List of processes (optional, used for concurrent execution)

            Returns:
                    None
    '''
    campaign_path = os.path.join(campaign_folder, "campaign")
    if not os.path.exists(campaign_path):
        os.makedirs(campaign_path)

    if concurrent_execution:
        process = Process(target=__write_test_table, args=(test_data, campaign_name, campaign_path))
        process.start()
        processes.append(process)
    else:
        __write_test_table(test_data, campaign_name, campaign_path)

def write_query_table(test_data: list, campaign_folder: str, processes: list=None) -> None:
    '''
    Creates a table with all query messages and their responses and timestamps. Saves the table as
    a CSV file and as an Excel file (optional).

            Parameters:
                    test_data (list): List of parsed test scenarios
                    campaign_folder (str): Path to the folder where the table should be saved
                    processes (list): List of processes (optional, used for concurrent execution)

            Returns:
                    None
    '''
    scenario_path = os.path.join(campaign_folder, "scenario")
    if not os.path.exists(scenario_path):
        os.makedirs(scenario_path)

    for test_scenario in test_data:
        if test_scenario[3] is not None:
            scenario_subpath = os.path.join(scenario_path, f"{test_scenario[0]['metadata']['t_uid']}")

            if concurrent_execution:
                process = Process(target=__write_query_table, args=(test_scenario, scenario_subpath))
                process.start()
                processes.append(process)
            else:
                __write_query_table(test_scenario, scenario_subpath)



def __write_test_table(test_data: list, campaign_name: str, output_path: str) -> None:
    filename = os.path.join(output_path, f"{campaign_name}_campaign_overview.csv")

    with open(filename, 'w', newline='', encoding='utf-8') as csvfile:
        writer = csv.writer(csvfile)

        # Write header
        writer.writerow(['Duration (s)', 'Method', 'Test-ID',
                         'Client', 'Server', 'Port',
                         'Cycle Time (ns)', 'Datagram Size (B)', 'QoS',
                         'Stress', 'Intensity', 'Location',
                         'Status',
                         'Losses [total]', 'Losses [ratio](%)', 'Losses [location]',
                         'Pakets [total]', 'PPS [udp]', 'PPS [ip]', 'Bandwidth [net](Mbps)', 'Bandwidth (gross)[Mbps]',
                         'Timer Misses',
                         'Remarks'])

        # Write content
        for scenario in test_data:
            scenario_data = __get_scenario_data(scenario)
            writer.writerow(scenario_data)

    __save_as_excel(filename, output_path)

def __get_scenario_data(scenario: tuple) -> list:
    scenario_description = scenario[0]
    scenario_client_results = scenario[1]
    scenario_server_results = scenario[2]

    scenario_data = list()

    # DURATION (Check if precise duration measurement was used)
    if scenario_client_results['report']['duration'] == -1:
        duration = scenario_description['duration']
    else:
        duration = scenario_client_results['report']['duration']
    scenario_data.append(duration)

    # METADATA
    scenario_data.append(scenario_description['metadata']['method'])
    scenario_data.append(scenario_description['metadata']['t_uid'])

    # CONNECTION
    scenario_data.append(scenario_description['connection']['client_ip'])
    scenario_data.append(scenario_description['connection']['server_ip'])
    scenario_data.append(scenario_description['connection']['port'])

    # TEST PARAMETERS
    scenario_data.append(scenario_description['connection']['cycle_time'])
    scenario_data.append(scenario_description['connection']['datagram_size'])
    scenario_data.append(scenario_description['connection']['qos'])
    scenario_data.append(scenario_description['stress']['type'])
    scenario_data.append(scenario_description['stress']['intensity'])
    scenario_data.append(scenario_description['stress']['location'])

    # STATUS
    scenario_data.append(scenario_client_results['status'])

    # LOSSES
    losses = scenario_client_results['report']['losses']
    losses_ratio = (losses / scenario_client_results['report']['total']) * 100
    losses_location = __get_losses_location(scenario_client_results, scenario_server_results)

    scenario_data.append(losses)
    scenario_data.append(losses_ratio)
    scenario_data.append(losses_location)

    # PACKETS
    pakets = scenario_client_results['report']['total']
    pps_udp = int(pakets // duration)
    pps_ip = pps_udp * (1 if scenario_description['connection']['datagram_size'] <= scenario_client_results['ip_statistic']['mtu'] else math.ceil(scenario_description['connection']['datagram_size'] / scenario_client_results['ip_statistic']['mtu']))
    bandwidth_net = pps_udp * scenario_description['connection']['datagram_size'] * 8 / 1000000
    if scenario_description['connection']['datagram_size'] <= scenario_client_results['ip_statistic']['mtu']:
        bandwidth_gross = pps_ip * (scenario_description['connection']['datagram_size'] + 42) * 8 / 1000000
    else:
        bandwidth_gross = pps_udp * (65000 + 280) * 8 / 1000000

    scenario_data.append(pakets)
    scenario_data.append(pps_udp)
    scenario_data.append(pps_ip)
    scenario_data.append(bandwidth_net)
    scenario_data.append(bandwidth_gross)

    # TIMER MISSES
    scenario_data.append(scenario_client_results['report']['timer_misses'])

    # REMARKS
    if scenario_server_results is None:
        scenario_data.append('Server data missing.')

    return scenario_data

def __get_losses_location(scenario_client_results: dict, scenario_server_results: dict) -> str:
    result = str()

    if scenario_client_results['report']['losses'] <= 0:
        return result

    # Check if server data exists
    server_data = scenario_server_results is not None

    # Check NIC statistics reported losses
    tx_dropped_client = scenario_client_results['ethtool_statistic']['tx_dropped']
    if tx_dropped_client > 0:
        result += 'Client (NIC)  '

    if not server_data:
        if tx_dropped_client == 0:
            result += 'Server [NO DATA]  Route (Switch)'
    else:
        tx_dropped_server = scenario_server_results['ethtool_statistic']['tx_dropped']
        if tx_dropped_server > 0:
            result += 'Server (NIC)  '

        udp_rec_err_server = scenario_server_results['netstat_statistic']['udp_rec_err']
        if udp_rec_err_server > 0:
            result += 'Server (UDP) [CAUSED BY STRESS?]  '

        if (tx_dropped_client == 0) and (tx_dropped_server == 0):
            result += 'Route (Switch)'

    return result

def __write_query_table(test_scenario: dict, output_path: str) -> None:
    if not os.path.exists(output_path):
        os.makedirs(output_path)
    filename = os.path.join(output_path, "query_overview.csv")

    duration = test_scenario[1]['report']['duration']
    losses = test_scenario[1]['report']['losses']
    total  = test_scenario[1]['report']['total']
    query = format_query(test_scenario[3], duration, losses, total)

    with open(filename, 'w', newline='', encoding='utf-8') as csvfile:
        writer = csv.writer(csvfile)

        # Write header
        writer.writerow(['Timestamp', 'Packets [total]', 'Losses [Total]', 'Losses [Difference]'])

        # Write content
        for report in query:
            writer.writerow([report['timestamp'], report['total'], report['losses'], report['difference']])

    __save_as_excel(filename, output_path)

def __save_as_excel(csv_path: str, output_path: str) -> None:
    if generate_excel:
        filename = os.path.join(output_path, f"{os.path.splitext(os.path.basename(csv_path))[0]}.xlsx")

        # Load the CSV data into a DataFrame
        df = pd.read_csv(csv_path)

        # Convert the DataFrame to an Excel file
        with pd.ExcelWriter(filename, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Sheet1', index=False)

            # Get the default sheet directly
            worksheet = writer.sheets['Sheet1']

            # Hide columns in campaign overview
            if "campaign_overview" in filename:
                for col in ['D', 'E', 'F']:
                    worksheet.column_dimensions[col].hidden = True

            # Create a table with all the data in the worksheet
            max_row = worksheet.max_row
            max_col = worksheet.max_column
            table = openpyxl.worksheet.table.Table(displayName="Table1", ref=f"A1:{chr(64 + max_col)}{max_row}")

            # Add the table to the worksheet
            worksheet.add_table(table)
